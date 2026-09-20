"""DOCX renderer (python-docx).

The layout is intentionally simple and ATS-friendly:
no images, no tables, headings are plain bold runs.
Template-specific behaviour is driven by ``template_config.load_template``.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt

from app.models.profile import Profile
from app.models.resume import TailoredResume
from app.resume.template_config import TemplateConfig, load_template


def render_docx(
    resume: TailoredResume,
    profile: Profile,
    output_path: Path,
    *,
    template: str | None = None,
) -> Path:
    cfg: TemplateConfig = load_template(template or resume.template)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    _render_header(doc, profile, resume)

    for section_name in cfg["section_order"]:
        if section_name == "summary":
            _render_summary(doc, resume)
        elif section_name == "skills":
            _render_skills(doc, resume)
        elif section_name == "experience":
            _render_experience(doc, resume, cfg)
        elif section_name == "projects":
            _render_projects(doc, resume, cfg)
        elif section_name == "education":
            _render_education(doc, resume)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def _section_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(13)


def _render_header(doc: Document, profile: Profile, resume: TailoredResume) -> None:
    name_p = doc.add_paragraph()
    name_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    name_run = name_p.add_run(profile.personal.full_name)
    name_run.bold = True
    name_run.font.size = Pt(18)

    contact_bits: list[str] = [profile.personal.email, profile.personal.phone]
    if profile.links.linkedin:
        contact_bits.append(str(profile.links.linkedin))
    if profile.links.github:
        contact_bits.append(str(profile.links.github))
    contact_p = doc.add_paragraph(" | ".join(contact_bits))
    contact_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    if resume.headline:
        headline_p = doc.add_paragraph(resume.headline)
        headline_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        headline_p.runs[0].italic = True


def _render_summary(doc: Document, resume: TailoredResume) -> None:
    _section_heading(doc, "Summary")
    doc.add_paragraph(resume.summary)


def _render_skills(doc: Document, resume: TailoredResume) -> None:
    if not resume.skills:
        return
    _section_heading(doc, "Skills")
    doc.add_paragraph(" • ".join(resume.skills))


def _render_experience(doc: Document, resume: TailoredResume, cfg: TemplateConfig) -> None:
    if not resume.employment:
        return
    _section_heading(doc, "Experience")
    for emp in resume.employment:
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(f"{emp.title} — {emp.company}")
        title_run.bold = True
        sub_p = doc.add_paragraph()
        sub_text = f"{emp.location or ''} | {emp.start} – {emp.end or 'Present'}"
        sub_run = sub_p.add_run(sub_text.strip(" |"))
        sub_run.italic = True
        for bullet in emp.bullets[: cfg["max_bullets_per_role"]]:
            doc.add_paragraph(bullet, style="List Bullet")
        if cfg["include_tech_line"] and emp.tech:
            tech_p = doc.add_paragraph()
            tech_run = tech_p.add_run("Tech: " + ", ".join(emp.tech))
            tech_run.italic = True


def _render_projects(doc: Document, resume: TailoredResume, cfg: TemplateConfig) -> None:
    if not resume.projects:
        return
    _section_heading(doc, "Projects")
    for proj in resume.projects:
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(proj.name)
        title_run.bold = True
        doc.add_paragraph(proj.description)
        for bullet in proj.bullets[: cfg["max_bullets_per_role"]]:
            doc.add_paragraph(bullet, style="List Bullet")
        if cfg["include_tech_line"] and proj.tech:
            tech_p = doc.add_paragraph()
            tech_run = tech_p.add_run("Tech: " + ", ".join(proj.tech))
            tech_run.italic = True


def _render_education(doc: Document, resume: TailoredResume) -> None:
    if not resume.education:
        return
    _section_heading(doc, "Education")
    for edu in resume.education:
        p = doc.add_paragraph()
        run = p.add_run(f"{edu.degree}, {edu.institution}")
        run.bold = True
        years = ""
        if edu.start_year and edu.end_year:
            years = f"{edu.start_year} – {edu.end_year}"
        elif edu.end_year:
            years = str(edu.end_year)
        if years:
            sub = doc.add_paragraph(years)
            sub.runs[0].italic = True
